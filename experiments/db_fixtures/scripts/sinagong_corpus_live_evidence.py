"""Load and verify the 20-workbook Sinagong corpus in local fixture DBs.

This module is product-external test-environment plumbing. It may load source
rows into local PostgreSQL/MySQL fixture services, but generated metadata is
explicitly marked TEST_ONLY_SYNTHETIC_METADATA and must not be promoted into
product Semantic Pack truth.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
from datetime import UTC, datetime, date, time
import hashlib
import json
import os
from pathlib import Path
import re
from typing import Any, Iterable, Mapping, Sequence
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook

from .fixture_modes import TEST_ONLY_MARKER

DEFAULT_DATASET_ROOT = Path("docs/reference/test_datasets/sinagong_tableau_2026")
DEFAULT_OUTPUT = Path("reports/reality/sinagong_20_live_db_fixture_evidence.json")
POSTGRES_DSN_ENV = "SEMANTIC_CONTEXT_POSTGRES_FIXTURE_DSN"
MYSQL_DSN_ENV = "SEMANTIC_CONTEXT_MYSQL_FIXTURE_DSN"
FIXTURE_GATE_ENV = "SEMANTIC_CONTEXT_FIXTURE_DB"
POSTGRES_SCHEMA = "semantic_fixture_sinagong_20"
MYSQL_DATABASE = "semantic_fixture_mysql_scope_c"
MANIFEST_TABLE = "semantic_fixture_manifest"


@dataclass(frozen=True)
class SheetLoadPlan:
    dataset_index: int
    sheet_index: int
    source_path: str
    source_name: str
    sheet_name: str
    table_name: str
    columns: list[str]
    headers: list[str]
    row_count: int
    column_count: int
    table_comment: str
    column_comments: dict[str, str]
    source_checksum: str

    def to_safe_dict(self) -> dict[str, Any]:
        # Do not include source row values in evidence; metadata/header names are
        # safe structural context for fixture verification.
        return asdict(self)


@dataclass(frozen=True)
class BackendEvidence:
    backend: str
    attempted: bool
    status: str
    fallback_used: bool
    reasons: list[str]
    dataset_count: int
    table_count: int
    row_count: int
    metadata_verified: bool
    manifest_rows: int
    table_comments_verified: int
    column_comments_verified: int


@dataclass(frozen=True)
class CorpusEvidence:
    fixture_only: bool
    metadata_mode: str
    synthetic_truth_blocked: bool
    source_dataset_count: int
    source_sheet_count: int
    source_row_count: int
    postgres: BackendEvidence
    mysql: BackendEvidence
    datasets: list[dict[str, Any]]
    generated_at: str


def discover_workbook_plans(dataset_root: Path) -> tuple[list[SheetLoadPlan], list[dict[str, Any]]]:
    files = sorted(dataset_root.rglob("*.xlsx"))
    plans: list[SheetLoadPlan] = []
    datasets: list[dict[str, Any]] = []
    for dataset_index, path in enumerate(files, start=1):
        checksum = _file_checksum(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        dataset_tables: list[dict[str, Any]] = []
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                rows = sheet.iter_rows(values_only=True)
                raw_headers = next(rows, ()) or ()
                headers = [_safe_header(value, index) for index, value in enumerate(raw_headers, start=1)]
                if not headers:
                    headers = ["_empty"]
                columns = [f"c{index:03d}" for index in range(1, len(headers) + 1)]
                row_count = sum(1 for row in rows if any(value is not None for value in row))
                table_name = _table_name(dataset_index, sheet_index, path, sheet.title)
                rel_path = str(path.relative_to(dataset_root))
                table_comment = (
                    f"{TEST_ONLY_MARKER}: fixture table for Sinagong dataset {dataset_index}; "
                    f"source={rel_path}; sheet={sheet.title}"
                )
                column_comments = {
                    column: f"{TEST_ONLY_MARKER}: source column header={header}; source={rel_path}; sheet={sheet.title}"
                    for column, header in zip(columns, headers, strict=True)
                }
                plan = SheetLoadPlan(
                    dataset_index=dataset_index,
                    sheet_index=sheet_index,
                    source_path=rel_path,
                    source_name=path.name,
                    sheet_name=sheet.title,
                    table_name=table_name,
                    columns=columns,
                    headers=headers,
                    row_count=row_count,
                    column_count=len(columns),
                    table_comment=table_comment,
                    column_comments=column_comments,
                    source_checksum=checksum,
                )
                plans.append(plan)
                dataset_tables.append({
                    "sheet_name": sheet.title,
                    "table_name": table_name,
                    "row_count": row_count,
                    "column_count": len(columns),
                })
        finally:
            workbook.close()
        datasets.append({
            "dataset_index": dataset_index,
            "source_path": str(path.relative_to(dataset_root)),
            "source_checksum": checksum,
            "sheet_count": len(dataset_tables),
            "tables": dataset_tables,
        })
    return plans, datasets


def load_postgres_corpus(*, dsn: str, plans: Sequence[SheetLoadPlan], dataset_root: Path) -> BackendEvidence:
    safety = _fixture_safety(dsn=dsn, expected_scheme={"postgres", "postgresql"}, backend="postgres")
    if safety:
        return _blocked_backend("postgres", safety, plans)
    try:
        import psycopg
    except ModuleNotFoundError as exc:
        return _blocked_backend("postgres", [f"psycopg is not installed: {exc}; no fallback backend was used"], plans)
    try:
        with psycopg.connect(dsn) as conn:
            with conn.cursor() as cur:
                cur.execute(f'CREATE SCHEMA IF NOT EXISTS {_pg_ident(POSTGRES_SCHEMA)}')
                _create_postgres_manifest(cur)
                for plan in plans:
                    _load_postgres_sheet(cur, plan, dataset_root)
                conn.commit()
            verification = _verify_postgres(conn, plans)
    except Exception as exc:  # pragma: no cover - live environment dependent.
        return _blocked_backend("postgres", [f"postgres live corpus load/verify failed explicitly: {exc}"], plans, attempted=True)
    return BackendEvidence(
        backend="postgres",
        attempted=True,
        status="loaded",
        fallback_used=False,
        reasons=[],
        dataset_count=len({plan.dataset_index for plan in plans}),
        table_count=len(plans),
        row_count=sum(plan.row_count for plan in plans),
        metadata_verified=verification["metadata_verified"],
        manifest_rows=verification["manifest_rows"],
        table_comments_verified=verification["table_comments_verified"],
        column_comments_verified=verification["column_comments_verified"],
    )


def load_mysql_corpus(*, dsn: str, plans: Sequence[SheetLoadPlan], dataset_root: Path) -> BackendEvidence:
    safety = _fixture_safety(dsn=dsn, expected_scheme={"mysql", "mysql+pymysql"}, backend="mysql")
    if safety:
        return _blocked_backend("mysql", safety, plans)
    try:
        import pymysql
    except ModuleNotFoundError as exc:
        return _blocked_backend("mysql", [f"pymysql is not installed: {exc}; no fallback backend was used"], plans)
    parsed = urlparse(dsn)
    database = unquote((parsed.path or "").lstrip("/")) or MYSQL_DATABASE
    try:
        conn = pymysql.connect(
            host=parsed.hostname or "localhost",
            port=parsed.port or 3306,
            user=unquote(parsed.username or ""),
            password=unquote(parsed.password or ""),
            database=database,
            charset="utf8mb4",
            connect_timeout=5,
            autocommit=False,
        )
        try:
            with conn.cursor() as cur:
                _create_mysql_manifest(cur)
                for plan in plans:
                    _load_mysql_sheet(cur, plan, dataset_root)
                conn.commit()
            verification = _verify_mysql(conn, plans, database)
        finally:
            conn.close()
    except Exception as exc:  # pragma: no cover - live environment dependent.
        return _blocked_backend("mysql", [f"mysql live corpus load/verify failed explicitly: {exc}"], plans, attempted=True)
    return BackendEvidence(
        backend="mysql",
        attempted=True,
        status="loaded",
        fallback_used=False,
        reasons=[],
        dataset_count=len({plan.dataset_index for plan in plans}),
        table_count=len(plans),
        row_count=sum(plan.row_count for plan in plans),
        metadata_verified=verification["metadata_verified"],
        manifest_rows=verification["manifest_rows"],
        table_comments_verified=verification["table_comments_verified"],
        column_comments_verified=verification["column_comments_verified"],
    )


def collect_corpus_evidence(
    *,
    dataset_root: Path = DEFAULT_DATASET_ROOT,
    postgres_dsn: str | None = None,
    mysql_dsn: str | None = None,
    env: Mapping[str, str] | None = None,
) -> CorpusEvidence:
    source_env = os.environ if env is None else env
    plans, datasets = discover_workbook_plans(dataset_root)
    pg_dsn = postgres_dsn or source_env.get(POSTGRES_DSN_ENV, "")
    my_dsn = mysql_dsn or source_env.get(MYSQL_DSN_ENV, "")
    postgres = (
        load_postgres_corpus(dsn=pg_dsn, plans=plans, dataset_root=dataset_root)
        if pg_dsn
        else _blocked_backend("postgres", [f"{POSTGRES_DSN_ENV} is not set; no fallback backend was used"], plans)
    )
    mysql = (
        load_mysql_corpus(dsn=my_dsn, plans=plans, dataset_root=dataset_root)
        if my_dsn
        else _blocked_backend("mysql", [f"{MYSQL_DSN_ENV} is not set; no fallback backend was used"], plans)
    )
    return CorpusEvidence(
        fixture_only=True,
        metadata_mode="test_only_synthetic_comments",
        synthetic_truth_blocked=True,
        source_dataset_count=len({plan.dataset_index for plan in plans}),
        source_sheet_count=len(plans),
        source_row_count=sum(plan.row_count for plan in plans),
        postgres=postgres,
        mysql=mysql,
        datasets=datasets,
        generated_at=datetime.now(UTC).isoformat(),
    )


def write_corpus_evidence(path: str | Path, evidence: CorpusEvidence) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(_json_safe(asdict(evidence)), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output


def _create_postgres_manifest(cur: Any) -> None:
    q_schema = _pg_ident(POSTGRES_SCHEMA)
    q_table = _pg_ident(MANIFEST_TABLE)
    cur.execute(f"DROP TABLE IF EXISTS {q_schema}.{q_table}")
    cur.execute(
        f"CREATE TABLE {q_schema}.{q_table} ("
        "dataset_index TEXT, source_path TEXT, sheet_name TEXT, table_name TEXT, "
        "row_count TEXT, column_count TEXT, source_checksum TEXT, metadata_mode TEXT, is_test_only TEXT)"
    )
    cur.execute(
        f"COMMENT ON TABLE {q_schema}.{q_table} IS {_pg_literal(f'{TEST_ONLY_MARKER}: manifest for local Sinagong 20-workbook fixture load')}"
    )


def _create_mysql_manifest(cur: Any) -> None:
    q_table = _my_ident(MANIFEST_TABLE)
    cur.execute(f"DROP TABLE IF EXISTS {q_table}")
    cur.execute(
        f"CREATE TABLE {q_table} ("
        "dataset_index TEXT, source_path TEXT, sheet_name TEXT, table_name TEXT, "
        "row_count TEXT, column_count TEXT, source_checksum TEXT, metadata_mode TEXT, is_test_only TEXT) "
        f"COMMENT={_my_literal(f'{TEST_ONLY_MARKER}: manifest for local Sinagong 20-workbook fixture load')}"
    )


def _load_postgres_sheet(cur: Any, plan: SheetLoadPlan, dataset_root: Path) -> None:
    q_schema = _pg_ident(POSTGRES_SCHEMA)
    q_table = _pg_ident(plan.table_name)
    cur.execute(f"DROP TABLE IF EXISTS {q_schema}.{q_table}")
    column_defs = ", ".join(f"{_pg_ident(column)} TEXT" for column in plan.columns)
    cur.execute(f"CREATE TABLE {q_schema}.{q_table} ({column_defs})")
    cur.execute(f"COMMENT ON TABLE {q_schema}.{q_table} IS {_pg_literal(plan.table_comment)}")
    for column, comment in plan.column_comments.items():
        cur.execute(f"COMMENT ON COLUMN {q_schema}.{q_table}.{_pg_ident(column)} IS {_pg_literal(comment)}")
    rows = list(_iter_sheet_rows(dataset_root / plan.source_path, plan.sheet_name, len(plan.columns)))
    if rows:
        placeholders = ", ".join(["%s"] * len(plan.columns))
        columns = ", ".join(_pg_ident(column) for column in plan.columns)
        cur.executemany(f"INSERT INTO {q_schema}.{q_table} ({columns}) VALUES ({placeholders})", rows)
    cur.execute(
        f"INSERT INTO {q_schema}.{_pg_ident(MANIFEST_TABLE)} VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (
            str(plan.dataset_index),
            plan.source_path,
            plan.sheet_name,
            plan.table_name,
            str(plan.row_count),
            str(plan.column_count),
            plan.source_checksum,
            "test_only_synthetic_comments",
            "true",
        ),
    )


def _load_mysql_sheet(cur: Any, plan: SheetLoadPlan, dataset_root: Path) -> None:
    q_table = _my_ident(plan.table_name)
    cur.execute(f"DROP TABLE IF EXISTS {q_table}")
    column_defs = ", ".join(
        f"{_my_ident(column)} TEXT COMMENT {_my_literal(plan.column_comments[column])}" for column in plan.columns
    )
    cur.execute(f"CREATE TABLE {q_table} ({column_defs}) COMMENT={_my_literal(plan.table_comment)}")
    rows = list(_iter_sheet_rows(dataset_root / plan.source_path, plan.sheet_name, len(plan.columns)))
    if rows:
        placeholders = ", ".join(["%s"] * len(plan.columns))
        columns = ", ".join(_my_ident(column) for column in plan.columns)
        cur.executemany(f"INSERT INTO {q_table} ({columns}) VALUES ({placeholders})", rows)
    cur.execute(
        f"INSERT INTO {_my_ident(MANIFEST_TABLE)} VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (
            str(plan.dataset_index),
            plan.source_path,
            plan.sheet_name,
            plan.table_name,
            str(plan.row_count),
            str(plan.column_count),
            plan.source_checksum,
            "test_only_synthetic_comments",
            "true",
        ),
    )


def _verify_postgres(conn: Any, plans: Sequence[SheetLoadPlan]) -> dict[str, Any]:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT COUNT(*) FROM information_schema.tables WHERE table_schema = %s AND table_name <> %s",
            (POSTGRES_SCHEMA, MANIFEST_TABLE),
        )
        table_count = int(cur.fetchone()[0])
        cur.execute(f"SELECT COUNT(*) FROM {_pg_ident(POSTGRES_SCHEMA)}.{_pg_ident(MANIFEST_TABLE)}")
        manifest_rows = int(cur.fetchone()[0])
        cur.execute(
            """
            SELECT COUNT(*)
            FROM pg_class c
            JOIN pg_namespace n ON n.oid = c.relnamespace
            WHERE n.nspname = %s
              AND c.relkind = 'r'
              AND c.relname <> %s
              AND obj_description(c.oid, 'pg_class') LIKE %s
            """,
            (POSTGRES_SCHEMA, MANIFEST_TABLE, f"%{TEST_ONLY_MARKER}%"),
        )
        table_comments = int(cur.fetchone()[0])
        cur.execute(
            """
            SELECT COUNT(*)
            FROM pg_attribute a
            JOIN pg_class c ON c.oid = a.attrelid
            JOIN pg_namespace n ON n.oid = c.relnamespace
            WHERE n.nspname = %s
              AND c.relkind = 'r'
              AND c.relname <> %s
              AND a.attnum > 0
              AND NOT a.attisdropped
              AND col_description(c.oid, a.attnum) LIKE %s
            """,
            (POSTGRES_SCHEMA, MANIFEST_TABLE, f"%{TEST_ONLY_MARKER}%"),
        )
        column_comments = int(cur.fetchone()[0])
    expected_columns = sum(plan.column_count for plan in plans)
    return {
        "manifest_rows": manifest_rows,
        "table_comments_verified": table_comments,
        "column_comments_verified": column_comments,
        "metadata_verified": table_count == len(plans) and manifest_rows == len(plans) and table_comments == len(plans) and column_comments == expected_columns,
    }


def _verify_mysql(conn: Any, plans: Sequence[SheetLoadPlan], database: str) -> dict[str, Any]:
    table_names = [plan.table_name for plan in plans]
    placeholders = ", ".join(["%s"] * len(table_names))
    with conn.cursor() as cur:
        cur.execute(
            f"SELECT COUNT(*) FROM information_schema.TABLES WHERE TABLE_SCHEMA = %s AND TABLE_NAME IN ({placeholders})",
            (database, *table_names),
        )
        table_count = int(cur.fetchone()[0])
        cur.execute(f"SELECT COUNT(*) FROM {_my_ident(MANIFEST_TABLE)}")
        manifest_rows = int(cur.fetchone()[0])
        cur.execute(
            f"""
            SELECT COUNT(*)
            FROM information_schema.TABLES
            WHERE TABLE_SCHEMA = %s
              AND TABLE_NAME IN ({placeholders})
              AND TABLE_COMMENT LIKE %s
            """,
            (database, *table_names, f"%{TEST_ONLY_MARKER}%"),
        )
        table_comments = int(cur.fetchone()[0])
        cur.execute(
            f"""
            SELECT COUNT(*)
            FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = %s
              AND TABLE_NAME IN ({placeholders})
              AND COLUMN_COMMENT LIKE %s
            """,
            (database, *table_names, f"%{TEST_ONLY_MARKER}%"),
        )
        column_comments = int(cur.fetchone()[0])
    expected_columns = sum(plan.column_count for plan in plans)
    return {
        "manifest_rows": manifest_rows,
        "table_comments_verified": table_comments,
        "column_comments_verified": column_comments,
        "metadata_verified": table_count == len(plans) and manifest_rows == len(plans) and table_comments == len(plans) and column_comments == expected_columns,
    }


def _iter_sheet_rows(path: Path, sheet_name: str, column_count: int) -> Iterable[tuple[str | None, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not any(value is not None for value in row):
                continue
            values = [_to_text(value) for value in list(row)[:column_count]]
            if len(values) < column_count:
                values.extend([None] * (column_count - len(values)))
            yield tuple(values)
    finally:
        workbook.close()


def _fixture_safety(*, dsn: str, expected_scheme: set[str], backend: str) -> list[str]:
    reasons: list[str] = []
    if os.environ.get(FIXTURE_GATE_ENV) != "1":
        reasons.append(f"{FIXTURE_GATE_ENV}=1 is required")
    parsed = urlparse(dsn)
    if parsed.scheme not in expected_scheme:
        reasons.append(f"only {backend} fixture DSNs are supported; no fallback backend was used")
    if (parsed.hostname or "") not in {"localhost", "127.0.0.1", "::1", "host.docker.internal"}:
        reasons.append("fixture DSN host must be local")
    database = (parsed.path or "").lstrip("/")
    if database and not database.startswith("semantic_fixture_"):
        reasons.append("database name should use semantic_fixture_* prefix for fixture safety")
    if any(token in dsn.casefold() for token in ("prod", "production", "warehouse", "analytics")):
        reasons.append("production-looking DSNs are refused for fixture loading")
    return reasons


def _blocked_backend(backend: str, reasons: list[str], plans: Sequence[SheetLoadPlan], *, attempted: bool = False) -> BackendEvidence:
    return BackendEvidence(
        backend=backend,
        attempted=attempted,
        status="blocked" if attempted else "pending",
        fallback_used=False,
        reasons=reasons,
        dataset_count=len({plan.dataset_index for plan in plans}),
        table_count=0,
        row_count=0,
        metadata_verified=False,
        manifest_rows=0,
        table_comments_verified=0,
        column_comments_verified=0,
    )


def _file_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()[:16]


def _table_name(dataset_index: int, sheet_index: int, path: Path, sheet_name: str) -> str:
    digest = hashlib.sha1(f"{path}:{sheet_name}".encode("utf-8")).hexdigest()[:8]
    return f"ds{dataset_index:02d}_s{sheet_index:02d}_{digest}"


def _safe_header(value: Any, index: int) -> str:
    text = str(value).strip() if value is not None else ""
    return text or f"column_{index:03d}"


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _pg_ident(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _pg_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"

def _my_ident(value: str) -> str:
    return "`" + value.replace("`", "``") + "`"


def _my_literal(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def _json_safe(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _json_safe(child) for key, child in value.items()}
    if isinstance(value, list):
        return [_json_safe(child) for child in value]
    return value


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-root", default=str(DEFAULT_DATASET_ROOT))
    parser.add_argument("--postgres-dsn", default=None)
    parser.add_argument("--mysql-dsn", default=None)
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args(argv)
    evidence = collect_corpus_evidence(
        dataset_root=Path(args.dataset_root),
        postgres_dsn=args.postgres_dsn,
        mysql_dsn=args.mysql_dsn,
    )
    path = write_corpus_evidence(args.output, evidence)
    print(path)
    if evidence.source_dataset_count != 20:
        print(f"expected 20 source datasets, found {evidence.source_dataset_count}")
        return 2
    if evidence.postgres.status != "loaded" or not evidence.postgres.metadata_verified:
        print(f"postgres verification failed: {evidence.postgres}")
        return 3
    if evidence.mysql.status != "loaded" or not evidence.mysql.metadata_verified:
        print(f"mysql verification failed: {evidence.mysql}")
        return 4
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
